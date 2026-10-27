import csv
import hashlib
import io
import os
from datetime import datetime, timezone
from typing import List, Optional

from fastapi import Depends, FastAPI, HTTPException, Query, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session

import auth
import models
import schemas
from database import SessionLocal, engine, get_db

models.Base.metadata.create_all(bind=engine)


def _bootstrap_super_admin() -> None:
    """Create the first super admin from env vars, if set and not present.

    Set ADMIN_EMAIL and ADMIN_PASSWORD in the environment (Vercel dashboard
    or local shell) before the first boot. Never breaks boot: any failure
    is ignored so the app still starts.
    """
    email = auth.normalize_email(os.getenv("ADMIN_EMAIL", ""))
    password = os.getenv("ADMIN_PASSWORD", "")
    if not email or not password or not auth.valid_email(email):
        return
    db = SessionLocal()
    try:
        if db.query(models.User).filter(models.User.email == email).first():
            return
        db.add(
            models.User(
                email=email,
                password_hash=auth.hash_password(password),
                role="super_admin",
                full_name="System Owner",
            )
        )
        db.commit()
    except Exception:
        db.rollback()
    finally:
        db.close()


_bootstrap_super_admin()

OPENING_PERIOD = "OPENING-BALANCE"

app = FastAPI(
    title="Residential Utility Usage Tracker API",
    description="Cumulative calculation, meter readings, paid/unpaid, edit/delete with chain recalculation, reporting. Version 4.2 adds roles and secure authentication.",
    version="4.2.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =====================================================================
# Helpers
# =====================================================================
def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def norm_id(value: str) -> str:
    return value.strip().upper()


def get_property_or_404(db: Session, prop_id: str) -> models.Property:
    record = db.query(models.Property).filter(models.Property.property_id == prop_id).first()
    if not record:
        raise HTTPException(status_code=404, detail=f"Property '{prop_id}' not found.")
    return record


def get_log_or_404(db: Session, log_id: int) -> models.UsageLog:
    log = db.query(models.UsageLog).filter(models.UsageLog.log_id == log_id).first()
    if not log:
        raise HTTPException(status_code=404, detail=f"Usage entry #{log_id} not found.")
    return log


def recalculate_chain(db: Session, prop: models.Property) -> None:
    """
    Property ke saare logs ko log_id ascending order me dobara calculate karta hai:
    previous/new cumulative units, previous meter reading, aur property ke totals
    (units, billed, paid, last meter reading). Edit/delete/payment ke baad call karo.
    Reading galat ho to rollback karke 400 deta hai.
    """
    logs = (
        db.query(models.UsageLog)
        .filter(models.UsageLog.property_id == prop.property_id)
        .order_by(models.UsageLog.log_id.asc())
        .all()
    )

    running_units = 0.0
    running_reading = 0.0
    total_billed = 0.0
    total_paid = 0.0

    for log in logs:
        log.previous_total_units = running_units

        if log.billing_period == OPENING_PERIOD:
            if log.current_reading is not None:
                running_reading = log.current_reading
            log.previous_reading = None
        elif log.current_reading is not None:
            if log.current_reading <= running_reading:
                db.rollback()
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"Entry '{log.billing_period}' (#{log.log_id}): current reading "
                        f"{log.current_reading:,.1f} previous reading {running_reading:,.1f} se zyada honi chahiye."
                    ),
                )
            log.previous_reading = running_reading
            log.added_units = round(log.current_reading - running_reading, 4)
            running_reading = log.current_reading
        else:
            # Manual units entry: meter reading chain me change nahi hoti
            log.previous_reading = None

        log.new_total_units = round(running_units + log.added_units, 4)
        running_units = log.new_total_units

        total_billed += log.bill_amount
        if log.payment_status == "PAID":
            total_paid += log.bill_amount

    prop.cumulative_units = running_units
    prop.last_meter_reading = running_reading
    prop.cumulative_billed_amount = round(total_billed, 2)
    prop.cumulative_paid_amount = round(total_paid, 2)
    prop.last_updated = now_utc()


def unpaid_query(db: Session, block: Optional[str] = None, search: Optional[str] = None):
    q = (
        db.query(models.UsageLog, models.Property)
        .join(models.Property, models.Property.property_id == models.UsageLog.property_id)
        .filter(models.UsageLog.payment_status == "UNPAID")
    )
    if block and block != "All Blocks":
        q = q.filter(models.Property.address_block == block)
    if search and search.strip():
        like = f"%{search.strip()}%"
        q = q.filter(
            or_(
                models.Property.property_id.ilike(like),
                models.Property.owner_name.ilike(like),
                models.UsageLog.billing_period.ilike(like),
            )
        )
    return q


def fmt_dt(value: Optional[datetime]) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else ""


# =====================================================================
# Auth, accounts, roles
# =====================================================================
def verify_home_link(
    db: Session, home_number: str, exclude_user_id: Optional[int] = None
) -> models.Property:
    """Home number ko verify karta hai: home maujood ho aur kisi aur customer
    se linked na ho. Warna 404/400. Customer ka apna home dobara link karna
    (exclude_user_id) allowed hai."""
    prop = db.query(models.Property).filter(models.Property.property_id == norm_id(home_number)).first()
    if not prop:
        raise HTTPException(
            status_code=404,
            detail=f"Home '{norm_id(home_number)}' not found. Pehle home create karo ya sahi home number do.",
        )
    clash = db.query(models.CustomerProfile).filter(models.CustomerProfile.home_number == prop.property_id)
    if exclude_user_id is not None:
        clash = clash.filter(models.CustomerProfile.user_id != exclude_user_id)
    if clash.first():
        raise HTTPException(
            status_code=400,
            detail="Ye home pehle se kisi aur customer se linked hai. Admin se rabta karo.",
        )
    return prop


def user_payload(user: models.User) -> schemas.UserResponse:
    return schemas.UserResponse.model_validate(user)


def profile_payload(profile: models.CustomerProfile | None) -> schemas.CustomerProfileResponse | None:
    return schemas.CustomerProfileResponse.model_validate(profile) if profile else None


@app.post("/auth/register", response_model=schemas.LoginResponse, status_code=status.HTTP_201_CREATED)
def register_user(payload: schemas.UserRegister, db: Session = Depends(get_db)):
    email = auth.normalize_email(payload.email)
    if not auth.valid_email(email):
        raise HTTPException(status_code=400, detail="Sahi email address do.")
    if db.query(models.User).filter(models.User.email == email).first():
        raise HTTPException(status_code=400, detail="Is email se account pehle se maujood hai.")
    account_number = payload.account_number.strip()
    if db.query(models.CustomerProfile).filter(models.CustomerProfile.account_number == account_number).first():
        raise HTTPException(status_code=400, detail="Ye account number pehle se istemal ho raha hai.")
    prop = None
    if payload.home_number and payload.home_number.strip():
        prop = verify_home_link(db, payload.home_number)
    user = models.User(
        email=email,
        password_hash=auth.hash_password(payload.password),
        role="customer",
        full_name=payload.full_name,
        phone=payload.phone,
    )
    db.add(user)
    db.flush()
    profile = models.CustomerProfile(
        user_id=user.id,
        account_number=account_number,
        home_number=prop.property_id if prop else None,
        phone=payload.phone,
    )
    db.add(profile)
    db.commit()
    db.refresh(user)
    db.refresh(profile)
    token = auth.create_session(db, user)
    return schemas.LoginResponse(token=token, user=user_payload(user), profile=profile_payload(profile))


@app.post("/auth/login", response_model=schemas.LoginResponse)
def login_user(payload: schemas.UserLogin, db: Session = Depends(get_db)):
    email = auth.normalize_email(payload.email)
    user = db.query(models.User).filter(models.User.email == email).first()
    if user is None or not auth.verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Email ya password ghalat hai.")
    if not user.is_active:
        raise HTTPException(status_code=401, detail="Account inactive hai. Admin se rabta karo.")
    token = auth.create_session(db, user)
    profile = db.query(models.CustomerProfile).filter(models.CustomerProfile.user_id == user.id).first()
    return schemas.LoginResponse(token=token, user=user_payload(user), profile=profile_payload(profile))


@app.post("/auth/logout")
def logout_user(
    credentials: HTTPAuthorizationCredentials = Depends(auth.bearer),
    db: Session = Depends(get_db),
):
    if credentials is not None and credentials.credentials:
        auth.destroy_session(db, credentials.credentials)
    return {"message": "Logged out."}


@app.get("/auth/me")
def get_me(user: models.User = Depends(auth.get_current_user), db: Session = Depends(get_db)):
    profile = db.query(models.CustomerProfile).filter(models.CustomerProfile.user_id == user.id).first()
    return {
        "user": user_payload(user).model_dump(),
        "profile": profile_payload(profile).model_dump() if profile else None,
    }


@app.post("/auth/change-password")
def change_password(
    payload: schemas.PasswordChange,
    credentials: HTTPAuthorizationCredentials = Depends(auth.bearer),
    user: models.User = Depends(auth.get_current_user),
    db: Session = Depends(get_db),
):
    if not auth.verify_password(payload.current_password, user.password_hash):
        raise HTTPException(status_code=400, detail="Current password ghalat hai.")
    user.password_hash = auth.hash_password(payload.new_password)
    user.force_password_change = False
    db.flush()
    # Baaki saare sessions khatam, sirf current session rehne do.
    current_hash = (
        hashlib.sha256(credentials.credentials.encode("utf-8")).hexdigest()
        if credentials is not None and credentials.credentials
        else None
    )
    q = db.query(models.Session).filter(models.Session.user_id == user.id)
    if current_hash:
        q = q.filter(models.Session.token_hash != current_hash)
    q.delete()
    db.commit()
    return {"message": "Password change ho gaya."}


@app.post("/auth/forgot-password")
def forgot_password(payload: schemas.ForgotPasswordRequest):
    # Email sender configured nahi hai, is liye jhoota reset link nahi bhejte.
    # Reset admin ke zariye hota hai (temporary password + force change).
    # Jawab hamesha ek jaisa, taake accounts enumerate na hon.
    return {
        "message": (
            "Password reset admin ke zariye hota hai. Apne partner ya super admin "
            "se rabta karo, woh tumhare liye temporary password issue karenge."
        )
    }


# =====================================================================
# Customer self-service
# =====================================================================
@app.get("/me/dashboard", response_model=schemas.CustomerDashboardResponse)
def customer_dashboard(user: models.User = Depends(auth.require_customer), db: Session = Depends(get_db)):
    profile = auth.get_customer_profile_or_404(db, user)
    prop = None
    logs = []
    if profile.home_number:
        prop = db.query(models.Property).filter(models.Property.property_id == profile.home_number).first()
    if prop is not None:
        logs = (
            db.query(models.UsageLog)
            .filter(models.UsageLog.property_id == prop.property_id)
            .order_by(models.UsageLog.log_id.desc())
            .all()
        )
    return schemas.CustomerDashboardResponse(
        user=user_payload(user),
        profile=profile_payload(profile),
        property=schemas.PropertyResponse.model_validate(prop) if prop else None,
        logs=logs,
        payment_status=auth.customer_payment_status(prop),
        total_billed=round(prop.cumulative_billed_amount, 2) if prop else 0.0,
        total_paid=round(prop.cumulative_paid_amount, 2) if prop else 0.0,
        total_pending=round(max(0.0, prop.cumulative_billed_amount - prop.cumulative_paid_amount), 2)
        if prop
        else 0.0,
    )


@app.put("/me/profile")
def update_my_profile(
    payload: schemas.ProfileUpdate,
    user: models.User = Depends(auth.require_customer),
    db: Session = Depends(get_db),
):
    profile = auth.get_customer_profile_or_404(db, user)
    if payload.full_name is not None:
        user.full_name = payload.full_name
    if payload.phone is not None:
        user.phone = payload.phone
        profile.phone = payload.phone
    if payload.address is not None:
        profile.address = payload.address
    profile.updated_at = now_utc()
    db.commit()
    db.refresh(user)
    db.refresh(profile)
    return {
        "user": user_payload(user).model_dump(),
        "profile": profile_payload(profile).model_dump(),
    }


@app.post("/me/link-home")
def link_my_home(
    payload: schemas.LinkHomeRequest,
    user: models.User = Depends(auth.require_customer),
    db: Session = Depends(get_db),
):
    profile = auth.get_customer_profile_or_404(db, user)
    prop = verify_home_link(db, payload.home_number, exclude_user_id=user.id)
    profile.home_number = prop.property_id
    profile.updated_at = now_utc()
    db.commit()
    db.refresh(profile)
    return {
        "message": f"Home {prop.property_id} link ho gaya.",
        "profile": profile_payload(profile).model_dump(),
    }


# =====================================================================
# Staff: customers list (partner + super admin)
# =====================================================================
@app.get("/customers/", response_model=List[schemas.CustomerListItem])
def list_customers(user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    rows = (
        db.query(models.CustomerProfile, models.User)
        .join(models.User, models.User.id == models.CustomerProfile.user_id)
        .order_by(models.CustomerProfile.account_number)
        .all()
    )
    items = []
    for profile, u in rows:
        prop = None
        if profile.home_number:
            prop = db.query(models.Property).filter(models.Property.property_id == profile.home_number).first()
        items.append(
            schemas.CustomerListItem(
                user_id=u.id,
                email=u.email,
                full_name=u.full_name,
                phone=u.phone,
                is_active=u.is_active,
                account_number=profile.account_number,
                home_number=profile.home_number,
                payment_status=auth.customer_payment_status(prop),
                total_billed=round(prop.cumulative_billed_amount, 2) if prop else 0.0,
                total_paid=round(prop.cumulative_paid_amount, 2) if prop else 0.0,
                total_pending=round(max(0.0, prop.cumulative_billed_amount - prop.cumulative_paid_amount), 2)
                if prop
                else 0.0,
            )
        )
    return items


# =====================================================================
# Super admin: user management
# =====================================================================
@app.post("/admin/users", response_model=schemas.UserResponse, status_code=status.HTTP_201_CREATED)
def admin_create_user(
    payload: schemas.AdminUserCreate,
    admin: models.User = Depends(auth.require_admin),
    db: Session = Depends(get_db),
):
    email = auth.normalize_email(payload.email)
    if not auth.valid_email(email):
        raise HTTPException(status_code=400, detail="Sahi email address do.")
    if payload.role not in auth.ROLES:
        raise HTTPException(status_code=400, detail="Role customer, partner ya super_admin honi chahiye.")
    if db.query(models.User).filter(models.User.email == email).first():
        raise HTTPException(status_code=400, detail="Is email se account pehle se maujood hai.")
    account_number = (payload.account_number or "").strip()
    if payload.role == "customer" and not account_number:
        raise HTTPException(status_code=400, detail="Customer ke liye account_number zaroori hai.")
    if account_number and db.query(models.CustomerProfile).filter(
        models.CustomerProfile.account_number == account_number
    ).first():
        raise HTTPException(status_code=400, detail="Ye account number pehle se istemal ho raha hai.")
    prop = None
    if payload.home_number and payload.home_number.strip():
        prop = verify_home_link(db, payload.home_number)
    user = models.User(
        email=email,
        password_hash=auth.hash_password(payload.password),
        role=payload.role,
        full_name=payload.full_name,
        phone=payload.phone,
    )
    db.add(user)
    db.flush()
    if payload.role == "customer":
        db.add(
            models.CustomerProfile(
                user_id=user.id,
                account_number=account_number,
                home_number=prop.property_id if prop else None,
                phone=payload.phone,
            )
        )
    db.commit()
    db.refresh(user)
    return user_payload(user)


@app.get("/admin/users")
def admin_list_users(admin: models.User = Depends(auth.require_admin), db: Session = Depends(get_db)):
    users = db.query(models.User).order_by(models.User.id).all()
    out = []
    for u in users:
        profile = db.query(models.CustomerProfile).filter(models.CustomerProfile.user_id == u.id).first()
        out.append(
            {
                "user": user_payload(u).model_dump(),
                "profile": profile_payload(profile).model_dump() if profile else None,
            }
        )
    return out


@app.patch("/admin/users/{user_id}/role")
def admin_set_role(
    user_id: int,
    payload: schemas.AdminRoleUpdate,
    admin: models.User = Depends(auth.require_admin),
    db: Session = Depends(get_db),
):
    target = db.query(models.User).filter(models.User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="User not found.")
    if target.id == admin.id:
        raise HTTPException(status_code=400, detail="Apni role khud change nahi kar sakte.")
    if target.role == "super_admin" and payload.role != "super_admin":
        remaining = (
            db.query(models.User)
            .filter(models.User.role == "super_admin", models.User.id != target.id)
            .count()
        )
        if remaining == 0:
            raise HTTPException(status_code=400, detail="Aakhri super admin ki role change nahi ho sakti.")
    if payload.role == "customer" and (
        db.query(models.CustomerProfile).filter(models.CustomerProfile.user_id == target.id).first() is None
    ):
        raise HTTPException(
            status_code=400,
            detail="Customer banane ke liye account_number chahiye. Is user ko customer profile ke sath dobara create karo.",
        )
    target.role = payload.role
    db.commit()
    db.refresh(target)
    return {"user": user_payload(target).model_dump()}


@app.patch("/admin/users/{user_id}/active")
def admin_set_active(
    user_id: int,
    payload: schemas.AdminActiveUpdate,
    admin: models.User = Depends(auth.require_admin),
    db: Session = Depends(get_db),
):
    target = db.query(models.User).filter(models.User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="User not found.")
    if target.id == admin.id and not payload.is_active:
        raise HTTPException(status_code=400, detail="Apna account khud deactivate nahi kar sakte.")
    target.is_active = payload.is_active
    if not payload.is_active:
        db.query(models.Session).filter(models.Session.user_id == target.id).delete()
    db.commit()
    db.refresh(target)
    return {"user": user_payload(target).model_dump()}


@app.post("/admin/users/{user_id}/reset-password")
def admin_reset_password(
    user_id: int,
    payload: schemas.AdminPasswordReset,
    admin: models.User = Depends(auth.require_admin),
    db: Session = Depends(get_db),
):
    target = db.query(models.User).filter(models.User.id == user_id).first()
    if not target:
        raise HTTPException(status_code=404, detail="User not found.")
    target.password_hash = auth.hash_password(payload.new_password)
    target.force_password_change = True
    db.query(models.Session).filter(models.Session.user_id == target.id).delete()
    db.commit()
    return {
        "message": (
            f"{target.email} ka password reset ho gaya. User ko temporary password dein; "
            "login ke baad password change karna zaroori hoga."
        )
    }


# =====================================================================
# Properties
# =====================================================================
@app.post("/properties/", response_model=schemas.PropertyResponse, status_code=status.HTTP_201_CREATED)
def create_property(prop: schemas.PropertyCreate, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    prop_id = norm_id(prop.property_id)
    if db.query(models.Property).filter(models.Property.property_id == prop_id).first():
        raise HTTPException(status_code=400, detail=f"Property ID '{prop_id}' already exists.")

    new_prop = models.Property(
        property_id=prop_id,
        owner_name=prop.owner_name,
        address_block=prop.address_block,
        cumulative_units=prop.initial_units,
        last_meter_reading=prop.initial_meter_reading,
    )
    db.add(new_prop)
    db.flush()

    if prop.initial_units > 0 or prop.initial_meter_reading > 0:
        db.add(
            models.UsageLog(
                property_id=prop_id,
                billing_period=OPENING_PERIOD,
                previous_total_units=0.0,
                added_units=prop.initial_units,
                new_total_units=prop.initial_units,
                previous_reading=None,
                current_reading=prop.initial_meter_reading,
                bill_amount=0.0,
                payment_status="PAID",
                paid_date=now_utc(),
            )
        )

    db.commit()
    db.refresh(new_prop)
    return new_prop


@app.post("/properties/bulk-generate/", status_code=status.HTTP_201_CREATED)
def bulk_generate_properties(payload: schemas.BulkPropertyCreate, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    created_ids: List[str] = []
    skipped_ids: List[str] = []
    prefix = payload.prefix.strip().upper()
    existing_ids = {pid for (pid,) in db.query(models.Property.property_id).all()}

    for i in range(payload.start_number, payload.start_number + payload.count):
        prop_id = f"{prefix}{i:03d}"
        if prop_id in existing_ids:
            skipped_ids.append(prop_id)
            continue
        db.add(
            models.Property(
                property_id=prop_id,
                owner_name=f"Resident {i}",
                address_block=payload.address_block,
                cumulative_units=0.0,
            )
        )
        created_ids.append(prop_id)

    db.commit()
    return {
        "message": f"Successfully generated {len(created_ids)} properties.",
        "created_count": len(created_ids),
        "skipped_count": len(skipped_ids),
        "sample_created": created_ids[:5],
    }


@app.put("/properties/{property_id}", response_model=schemas.PropertyResponse)
def update_property_details(property_id: str, payload: schemas.PropertyUpdate, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    property_record = get_property_or_404(db, norm_id(property_id))

    if payload.owner_name is not None:
        property_record.owner_name = payload.owner_name
    if payload.address_block is not None:
        property_record.address_block = payload.address_block

    property_record.last_updated = now_utc()
    db.commit()
    db.refresh(property_record)
    return property_record


@app.put("/properties/{property_id}/rename", response_model=schemas.PropertyResponse)
def rename_property_id(property_id: str, payload: schemas.PropertyRename, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    old_id = norm_id(property_id)
    new_id = norm_id(payload.new_property_id)

    if not new_id:
        raise HTTPException(status_code=400, detail="New Home ID cannot be empty.")
    if new_id == old_id:
        raise HTTPException(status_code=400, detail="New Home ID is same as the current one.")

    get_property_or_404(db, old_id)

    if db.query(models.Property).filter(models.Property.property_id == new_id).first():
        raise HTTPException(status_code=400, detail=f"Property ID '{new_id}' already exists.")

    # Postgres: ON UPDATE CASCADE logs ko khud move kar deta hai.
    # SQLite: foreign keys enforce nahi hoti, is liye logs ko neeche manually update karte hain.
    db.query(models.Property).filter(models.Property.property_id == old_id).update(
        {"property_id": new_id, "last_updated": now_utc()}, synchronize_session=False
    )
    db.query(models.UsageLog).filter(models.UsageLog.property_id == old_id).update(
        {"property_id": new_id}, synchronize_session=False
    )
    db.commit()

    return db.query(models.Property).filter(models.Property.property_id == new_id).first()


@app.get("/properties/", response_model=List[schemas.PropertyResponse])
def get_all_properties(skip: int = 0, limit: int = 2000, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    return db.query(models.Property).order_by(models.Property.property_id).offset(skip).limit(limit).all()


@app.get("/properties/{property_id}", response_model=schemas.PropertyDetailResponse)
def get_property_history(property_id: str, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    return get_property_or_404(db, norm_id(property_id))


@app.post("/properties/{property_id}/pay-all")
def pay_all_for_property(property_id: str, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    prop = get_property_or_404(db, norm_id(property_id))

    unpaid_logs = (
        db.query(models.UsageLog)
        .filter(models.UsageLog.property_id == prop.property_id, models.UsageLog.payment_status == "UNPAID")
        .all()
    )
    if not unpaid_logs:
        raise HTTPException(status_code=400, detail=f"'{prop.property_id}' ki koi unpaid entry nahi hai.")

    paid_time = now_utc()
    amount = 0.0
    for log in unpaid_logs:
        log.payment_status = "PAID"
        log.paid_date = paid_time
        amount += log.bill_amount

    db.flush()
    recalculate_chain(db, prop)
    db.commit()
    return {
        "message": f"{len(unpaid_logs)} entries paid mark ho gayin for {prop.property_id}.",
        "entries_paid": len(unpaid_logs),
        "amount_paid": round(amount, 2),
    }


# =====================================================================
# Usage entries
# =====================================================================
@app.post("/usage/add/", response_model=schemas.UsageLogResponse, status_code=status.HTTP_201_CREATED)
def add_usage_entry(entry: schemas.UsageEntryCreate, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    prop_id = norm_id(entry.property_id)
    prop = get_property_or_404(db, prop_id)

    previous_units = prop.cumulative_units

    if entry.current_reading is not None:
        previous_reading = prop.last_meter_reading
        if entry.current_reading <= previous_reading:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Current reading ({entry.current_reading:,.1f}) previous reading "
                    f"({previous_reading:,.1f}) se zyada honi chahiye."
                ),
            )
        units = round(entry.current_reading - previous_reading, 4)
        current_reading = entry.current_reading
    else:
        previous_reading = None
        current_reading = None
        units = entry.added_units

    usage_log = models.UsageLog(
        property_id=prop_id,
        billing_period=entry.billing_period.strip(),
        previous_total_units=previous_units,
        added_units=units,
        new_total_units=round(previous_units + units, 4),
        previous_reading=previous_reading,
        current_reading=current_reading,
        bill_amount=entry.bill_amount,
        payment_status="UNPAID",
    )
    db.add(usage_log)
    db.flush()
    recalculate_chain(db, prop)
    db.commit()
    db.refresh(usage_log)
    return usage_log


@app.put("/usage/{log_id}", response_model=schemas.UsageLogResponse)
def update_usage_entry(log_id: int, payload: schemas.UsageEntryUpdate, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    log = get_log_or_404(db, log_id)
    if log.billing_period == OPENING_PERIOD:
        raise HTTPException(status_code=400, detail="Opening balance entry edit nahi ho sakti.")

    prop = get_property_or_404(db, log.property_id)

    if payload.billing_period is not None:
        log.billing_period = payload.billing_period.strip()
    if payload.bill_amount is not None:
        log.bill_amount = payload.bill_amount

    if payload.current_reading is not None:
        log.current_reading = payload.current_reading
    elif payload.added_units is not None:
        log.added_units = payload.added_units
        log.current_reading = None
        log.previous_reading = None

    log.updated_at = now_utc()
    db.flush()
    recalculate_chain(db, prop)
    db.commit()
    db.refresh(log)
    return log


@app.delete("/usage/undo-last/{property_id}")
def undo_last_usage_entry(property_id: str, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    prop = get_property_or_404(db, norm_id(property_id))

    last_log = (
        db.query(models.UsageLog)
        .filter(models.UsageLog.property_id == prop.property_id)
        .order_by(models.UsageLog.log_id.desc())
        .first()
    )
    if not last_log:
        raise HTTPException(status_code=400, detail=f"No usage history found to undo for '{prop.property_id}'.")
    if last_log.billing_period == OPENING_PERIOD:
        raise HTTPException(status_code=400, detail="Sirf opening balance bacha hai, usay undo nahi kiya ja sakta.")

    reverted_info = {"undone_log_id": last_log.log_id, "removed_units": last_log.added_units}

    db.delete(last_log)
    db.flush()
    recalculate_chain(db, prop)
    db.commit()

    reverted_info["restored_cumulative_units"] = prop.cumulative_units
    return {"message": f"Successfully reverted last entry for {prop.property_id}.", "details": reverted_info}


@app.get("/usage/unpaid/")
def get_unpaid_entries(
    block: Optional[str] = None,
    search: Optional[str] = None,
    skip: int = Query(default=0, ge=0),
    limit: int = Query(default=100, ge=1, le=2000),
    user: models.User = Depends(auth.require_staff),
    db: Session = Depends(get_db),
):
    q = unpaid_query(db, block, search)
    total_count = q.count()
    total_amount = q.with_entities(func.coalesce(func.sum(models.UsageLog.bill_amount), 0.0)).scalar() or 0.0

    rows = q.order_by(models.UsageLog.log_id.desc()).offset(skip).limit(limit).all()
    items = [
        {
            "log_id": log.log_id,
            "property_id": prop.property_id,
            "owner_name": prop.owner_name,
            "address_block": prop.address_block,
            "billing_period": log.billing_period,
            "added_units": log.added_units,
            "bill_amount": log.bill_amount,
            "entry_date": log.entry_date,
        }
        for log, prop in rows
    ]
    return {"total_count": total_count, "total_amount": round(float(total_amount), 2), "items": items}


@app.delete("/usage/{log_id}")
def delete_usage_entry(log_id: int, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    log = get_log_or_404(db, log_id)
    if log.billing_period == OPENING_PERIOD:
        raise HTTPException(status_code=400, detail="Opening balance entry delete nahi ho sakti.")

    prop = get_property_or_404(db, log.property_id)
    removed = {
        "deleted_log_id": log.log_id,
        "billing_period": log.billing_period,
        "removed_units": log.added_units,
        "removed_bill": log.bill_amount,
    }

    db.delete(log)
    db.flush()
    recalculate_chain(db, prop)
    db.commit()

    return {
        "message": f"Entry #{log_id} delete ho gayi. {prop.property_id} ke totals dobara calculate ho gaye.",
        "details": removed,
        "cumulative_units": prop.cumulative_units,
        "last_meter_reading": prop.last_meter_reading,
    }


@app.patch("/usage/{log_id}/payment", response_model=schemas.UsageLogResponse)
def update_payment_status(log_id: int, payload: schemas.PaymentUpdate, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    log = get_log_or_404(db, log_id)
    prop = get_property_or_404(db, log.property_id)

    log.payment_status = payload.status
    log.paid_date = now_utc() if payload.status == "PAID" else None

    db.flush()
    recalculate_chain(db, prop)
    db.commit()
    db.refresh(log)
    return log


@app.get("/logs/", response_model=List[schemas.UsageLogResponse])
def get_all_usage_logs(limit: int = 5000, user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    return db.query(models.UsageLog).order_by(models.UsageLog.log_id.desc()).limit(limit).all()


# =====================================================================
# Dashboard
# =====================================================================
@app.get("/dashboard/summary", response_model=schemas.DashboardSummary)
def dashboard_summary(user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    P = models.Property
    total_homes, active_homes, total_units, total_billed, total_paid = db.query(
        func.count(P.property_id),
        func.coalesce(func.sum(case((P.cumulative_units > 0, 1), else_=0)), 0),
        func.coalesce(func.sum(P.cumulative_units), 0.0),
        func.coalesce(func.sum(P.cumulative_billed_amount), 0.0),
        func.coalesce(func.sum(P.cumulative_paid_amount), 0.0),
    ).one()

    unpaid_count = (
        db.query(func.count(models.UsageLog.log_id))
        .filter(models.UsageLog.payment_status == "UNPAID")
        .scalar()
        or 0
    )

    return schemas.DashboardSummary(
        total_homes=int(total_homes or 0),
        active_homes=int(active_homes or 0),
        total_units=round(float(total_units), 2),
        total_billed=round(float(total_billed), 2),
        total_paid=round(float(total_paid), 2),
        total_pending=round(max(0.0, float(total_billed) - float(total_paid)), 2),
        unpaid_entries_count=int(unpaid_count),
    )


# =====================================================================
# Export (Excel / CSV) - sirf openpyxl + csv, taake Vercel bundle halka rahe
# =====================================================================
PROPS_HEADERS = [
    "Home ID", "Resident Name", "Block", "Total Cumulative Units", "Last Meter Reading",
    "Total Billed (Rs.)", "Total Paid (Rs.)", "Pending (Rs.)", "Last Updated",
]


def props_rows(db: Session):
    props = db.query(models.Property).order_by(models.Property.property_id).all()
    return [
        [
            p.property_id, p.owner_name, p.address_block, p.cumulative_units, p.last_meter_reading,
            p.cumulative_billed_amount, p.cumulative_paid_amount,
            round(max(0.0, p.cumulative_billed_amount - p.cumulative_paid_amount), 2),
            fmt_dt(p.last_updated),
        ]
        for p in props
    ]


def build_export_tables(db: Session):
    logs = db.query(models.UsageLog).order_by(models.UsageLog.log_id.desc()).all()
    logs_rows = [
        [
            l.log_id, l.property_id, l.billing_period, l.previous_reading, l.current_reading,
            l.previous_total_units, l.added_units, l.new_total_units, l.bill_amount,
            l.payment_status, fmt_dt(l.paid_date), fmt_dt(l.entry_date), fmt_dt(l.updated_at),
        ]
        for l in logs
    ]
    unpaid_rows = [
        [
            prop.property_id, prop.owner_name, prop.address_block, log.billing_period,
            log.added_units, log.bill_amount, fmt_dt(log.entry_date),
        ]
        for log, prop in unpaid_query(db).order_by(models.Property.property_id, models.UsageLog.log_id).all()
    ]
    return [
        ("Master_Homes_Summary", PROPS_HEADERS, props_rows(db)),
        (
            "All_Usage_History",
            [
                "Entry #", "Home ID", "Billing Period", "Previous Reading", "Current Reading",
                "Previous Total Units", "Added Units", "New Cumulative Total", "Bill Amount (Rs.)",
                "Status", "Paid Date", "Entry Date", "Edited At",
            ],
            logs_rows,
        ),
        (
            "Unpaid_Dues",
            ["Home ID", "Resident Name", "Block", "Billing Period", "Units", "Bill Amount (Rs.)", "Entry Date"],
            unpaid_rows,
        ),
    ]


@app.get("/export/excel")
def export_excel(user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in build_export_tables(db):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 2)
        ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Utility_Tracker_Report_{stamp}.xlsx"'},
    )


@app.get("/export/csv")
def export_csv(user: models.User = Depends(auth.require_staff), db: Session = Depends(get_db)):
    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow(PROPS_HEADERS)
    writer.writerows(props_rows(db))
    data = text.getvalue().encode("utf-8-sig")  # utf-8-sig: Excel me sahi khulta hai

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="Master_Properties_{stamp}.csv"'},
    )


# =====================================================================
# Frontend: public/index.html
# Local par yahan se serve hota hai. Vercel par public/ folder CDN khud serve karta hai.
# Mount sab routes ke BAAD hona zaroori hai.
# =====================================================================
PUBLIC_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "public")
if os.path.isdir(PUBLIC_DIR):
    app.mount("/", StaticFiles(directory=PUBLIC_DIR, html=True), name="public")
